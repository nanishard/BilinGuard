# -*- coding: utf-8 -*-
"""
Bootstrap 95% CI — loads from saved manifests (dbil/ibil) to ensure
exact same val split as training. Also evaluates eyelid TBIL and type models.
"""
import os, re, json, random, numpy as np, pandas as pd, cv2
from PIL import Image
from tqdm import tqdm
os.environ['HF_ENDPOINT'] = 'https://hf-mirror.com'
os.environ['HF_HUB_OFFLINE'] = '1'
import torch, torch.nn.functional as F
from torchvision import transforms
import timm
from sklearn.metrics import roc_auc_score, accuracy_score, f1_score
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt

BASE = r'D:\research\人脸识别营养\传染科'
MODEL = os.path.join(BASE, 'models')
RES = os.path.join(BASE, 'results')
FIG = os.path.join(RES, 'figures', 'bootstrap')
TBL = os.path.join(RES, 'tables')
os.makedirs(FIG, exist_ok=True)
DEV = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
SEED = 42; IMG_SIZE = 224; N_BOOT = 1000
random.seed(SEED); np.random.seed(SEED); torch.manual_seed(SEED)
plt.rcParams.update({'font.family': 'Arial', 'font.size': 9, 'axes.linewidth': 0.6,
                      'figure.dpi': 300, 'savefig.dpi': 300, 'savefig.bbox': 'tight', 'pdf.fonttype': 42})


def read_img(path):
    try:
        fb = np.fromfile(path, dtype=np.uint8)
        img = cv2.imdecode(fb, cv2.IMREAD_COLOR)
        return cv2.cvtColor(img, cv2.COLOR_BGR2RGB) if img is not None else None
    except Exception:
        return None

def clahe(img_rgb):
    lab = cv2.cvtColor(img_rgb, cv2.COLOR_RGB2LAB)
    l, a, b = cv2.split(lab)
    l = cv2.createCLAHE(clipLimit=3.0, tileGridSize=(8, 8)).apply(l)
    return cv2.cvtColor(cv2.merge([l, a, b]), cv2.COLOR_LAB2RGB)

ev_tf = transforms.Compose([
    transforms.Resize((IMG_SIZE, IMG_SIZE)), transforms.ToTensor(),
    transforms.Normalize([0.485, 0.456, 0.406], [0.229, 0.224, 0.225])])


def predict_patient(model, paths, n_max=12):
    valid = [p for p in paths[:n_max] if os.path.exists(p)]
    if not valid: return None
    tensors = []
    for p in valid:
        img = read_img(p)
        if img is None: continue
        img = clahe(img)
        tensors.append(ev_tf(Image.fromarray(img).convert('RGB')).unsqueeze(0))
    if not tensors: return None
    t = torch.cat(tensors).to(DEV)
    with torch.no_grad():
        return F.softmax(model(t), dim=1).mean(0).cpu().numpy()


def bootstrap_metrics(true, probs, nc, n_boot=N_BOOT):
    rng = np.random.RandomState(SEED)
    n = len(true)
    aucs, accs, f1s = [], [], []
    for _ in range(n_boot):
        idx = rng.randint(0, n, size=n)
        yt = true[idx]; yp = probs[idx]
        if len(np.unique(yt)) < 2: continue
        pred = yp.argmax(1)
        try:
            a = roc_auc_score(yt, yp[:, 1]) if nc == 2 else roc_auc_score(yt, yp, multi_class='ovr', labels=list(range(nc)))
            if not np.isnan(a): aucs.append(a)
        except Exception: pass
        accs.append(accuracy_score(yt, pred))
        try: f1s.append(f1_score(yt, pred, average='macro', labels=list(range(nc)), zero_division=0))
        except Exception: f1s.append(0)
    def pct(arr):
        arr = [x for x in arr if not np.isnan(x)]
        return (float(np.percentile(arr, 2.5)), float(np.percentile(arr, 97.5))) if arr else (0, 0)
    al, ah = pct(aucs); al2, ah2 = pct(accs); fl, fh = pct(f1s)
    return {'auc_lo': al, 'auc_hi': ah, 'acc_lo': al2, 'acc_hi': ah2, 'f1_lo': fl, 'f1_hi': fh,
            'n_boot_auc': len(aucs)}


def get_preds_and_ci(name, backbone_id, ckpt, nc, val_items, img_key):
    """Load model, predict, bootstrap. Returns (result_dict, probs, true)."""
    path = os.path.join(MODEL, ckpt)
    if not os.path.exists(path):
        print(f'    SKIP {name}: not found'); return None
    m = timm.create_model(backbone_id, pretrained=False, num_classes=nc).to(DEV)
    m.load_state_dict(torch.load(path, map_location=DEV, weights_only=True)); m.eval()
    probs, true = [], []
    for item in tqdm(val_items, desc=name, leave=False):
        imgs = item.get(img_key, item.get('faces', []))
        pr = predict_patient(m, imgs)
        if pr is None: pr = np.ones(nc)/nc
        probs.append(pr); true.append(item['label'])
    probs = np.array(probs); true = np.array(true)
    del m; torch.cuda.empty_cache()
    bs = bootstrap_metrics(true, probs, nc)
    try:
        pa = roc_auc_score(true, probs[:,1]) if nc == 2 else roc_auc_score(true, probs, multi_class='ovr', labels=list(range(nc)))
    except Exception: pa = 0
    res = {'name': name, 'nc': nc, 'n_val': len(true),
           'point_auc': pa, 'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
           'point_acc': accuracy_score(true, probs.argmax(1)),
           'point_f1': f1_score(true, probs.argmax(1), average='macro', zero_division=0)}
    print(f'    {name}: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]  (n_val={len(true)})')
    return res, probs, true


def split_from_manifest(mdf, grade_col, ratio=0.2):
    """Exact same split as training scripts."""
    by_grade = {}
    for _, r in mdf.iterrows():
        g = int(r[grade_col])
        by_grade.setdefault(g, []).append(r['patient_id'])
    tr, va = set(), set()
    for g, pids in by_grade.items():
        random.seed(SEED); random.shuffle(pids)
        n = max(1, int(len(pids) * ratio))
        va.update(pids[:n]); tr.update(pids[n:])
    return tr, va


BACKBONES = [('convnext', 'convnext_tiny'), ('vit', 'vit_tiny_patch16_224'),
             ('efficientnet', 'efficientnet_b0'), ('swin', 'swin_tiny_patch4_window7_224')]

all_results = []

# ═════════════════════════════════════════════════════════════
# DBIL — from dbil_manifest.csv
# ═════════════════════════════════════════════════════════════
print('='*60)
print('  DBIL Bootstrap')
print('='*60)
dbil_mdf = pd.read_csv(os.path.join(BASE, 'data', 'dbil_manifest.csv'))
tr_pids, va_pids = split_from_manifest(dbil_mdf, 'dbil_grade')
va_df = dbil_mdf[dbil_mdf['patient_id'].isin(va_pids)]

va_eye = [{'eyelids': json.loads(r['eyelid_images']), 'faces': json.loads(r['face_images']),
            'label': int(r['dbil_grade']), 'pid': r['patient_id']}
           for _, r in va_df.iterrows() if r['n_eyelid'] > 0 and json.loads(r['eyelid_images'])]
va_face = [{'faces': json.loads(r['face_images']), 'label': int(r['dbil_grade']), 'pid': r['patient_id']}
            for _, r in va_df.iterrows() if r['n_face'] > 0]

from collections import Counter
print(f'  Eyelid val: {len(va_eye)} {Counter(p["label"] for p in va_eye)}')
print(f'  Face val: {len(va_face)} {Counter(p["label"] for p in va_face)}')

print('\n  [DBIL Eyelid]')
dbil_eye_preds, dbil_eye_true = {}, None
for name, bb in BACKBONES:
    r = get_preds_and_ci(f'DBIL-Eyelid-{name}', bb, f'dbil_eyelid_{name}.pt', 3, va_eye, 'eyelids')
    if r:
        res, probs, true = r; res['task'] = 'DBIL'; res['modality'] = 'Eyelid'
        all_results.append(res); dbil_eye_preds[name] = probs; dbil_eye_true = true

if len(dbil_eye_preds) >= 2 and dbil_eye_true is not None:
    ens = np.mean(list(dbil_eye_preds.values()), axis=0)
    bs = bootstrap_metrics(dbil_eye_true, ens, 3)
    pa = roc_auc_score(dbil_eye_true, ens, multi_class='ovr', labels=[0,1,2])
    all_results.append({'name': 'DBIL-Eyelid-Ensemble', 'task': 'DBIL', 'modality': 'Eyelid',
                         'nc': 3, 'n_val': len(dbil_eye_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(dbil_eye_true, ens.argmax(1)),
                         'point_f1': f1_score(dbil_eye_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    DBIL-Eyelid-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

print('\n  [DBIL Face]')
dbil_face_preds, dbil_face_true = {}, None
for name, bb in BACKBONES:
    r = get_preds_and_ci(f'DBIL-Face-{name}', bb, f'dbil_face_{name}.pt', 3, va_face, 'faces')
    if r:
        res, probs, true = r; res['task'] = 'DBIL'; res['modality'] = 'Face'
        all_results.append(res); dbil_face_preds[name] = probs; dbil_face_true = true

if len(dbil_face_preds) >= 2 and dbil_face_true is not None:
    ens = np.mean(list(dbil_face_preds.values()), axis=0)
    bs = bootstrap_metrics(dbil_face_true, ens, 3)
    pa = roc_auc_score(dbil_face_true, ens, multi_class='ovr', labels=[0,1,2])
    all_results.append({'name': 'DBIL-Face-Ensemble', 'task': 'DBIL', 'modality': 'Face',
                         'nc': 3, 'n_val': len(dbil_face_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(dbil_face_true, ens.argmax(1)),
                         'point_f1': f1_score(dbil_face_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    DBIL-Face-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

# ═════════════════════════════════════════════════════════════
# IBIL — from ibil_manifest.csv
# ═════════════════════════════════════════════════════════════
print('\n' + '='*60)
print('  IBIL Bootstrap')
print('='*60)
ibil_mdf = pd.read_csv(os.path.join(BASE, 'data', 'ibil_manifest.csv'))
tr_pids_i, va_pids_i = split_from_manifest(ibil_mdf, 'ibil_grade')
va_df_i = ibil_mdf[ibil_mdf['patient_id'].isin(va_pids_i)]

va_eye_i = [{'eyelids': json.loads(r['eyelid_images']), 'faces': json.loads(r['face_images']),
              'label': int(r['ibil_grade']), 'pid': r['patient_id']}
             for _, r in va_df_i.iterrows() if r['n_eyelid'] > 0 and json.loads(r['eyelid_images'])]
va_face_i = [{'faces': json.loads(r['face_images']), 'label': int(r['ibil_grade']), 'pid': r['patient_id']}
              for _, r in va_df_i.iterrows() if r['n_face'] > 0]

print(f'  Eyelid val: {len(va_eye_i)} {Counter(p["label"] for p in va_eye_i)}')
print(f'  Face val: {len(va_face_i)} {Counter(p["label"] for p in va_face_i)}')

print('\n  [IBIL Eyelid]')
ibil_eye_preds, ibil_eye_true = {}, None
for name, bb in BACKBONES:
    r = get_preds_and_ci(f'IBIL-Eyelid-{name}', bb, f'ibil_eyelid_{name}.pt', 3, va_eye_i, 'eyelids')
    if r:
        res, probs, true = r; res['task'] = 'IBIL'; res['modality'] = 'Eyelid'
        all_results.append(res); ibil_eye_preds[name] = probs; ibil_eye_true = true

if len(ibil_eye_preds) >= 2 and ibil_eye_true is not None:
    ens = np.mean(list(ibil_eye_preds.values()), axis=0)
    bs = bootstrap_metrics(ibil_eye_true, ens, 3)
    pa = roc_auc_score(ibil_eye_true, ens, multi_class='ovr', labels=[0,1,2])
    all_results.append({'name': 'IBIL-Eyelid-Ensemble', 'task': 'IBIL', 'modality': 'Eyelid',
                         'nc': 3, 'n_val': len(ibil_eye_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(ibil_eye_true, ens.argmax(1)),
                         'point_f1': f1_score(ibil_eye_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    IBIL-Eyelid-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

print('\n  [IBIL Face]')
ibil_face_preds, ibil_face_true = {}, None
for name, bb in BACKBONES:
    r = get_preds_and_ci(f'IBIL-Face-{name}', bb, f'ibil_face_{name}.pt', 3, va_face_i, 'faces')
    if r:
        res, probs, true = r; res['task'] = 'IBIL'; res['modality'] = 'Face'
        all_results.append(res); ibil_face_preds[name] = probs; ibil_face_true = true

if len(ibil_face_preds) >= 2 and ibil_face_true is not None:
    ens = np.mean(list(ibil_face_preds.values()), axis=0)
    bs = bootstrap_metrics(ibil_face_true, ens, 3)
    pa = roc_auc_score(ibil_face_true, ens, multi_class='ovr', labels=[0,1,2])
    all_results.append({'name': 'IBIL-Face-Ensemble', 'task': 'IBIL', 'modality': 'Face',
                         'nc': 3, 'n_val': len(ibil_face_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(ibil_face_true, ens.argmax(1)),
                         'point_f1': f1_score(ibil_face_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    IBIL-Face-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

# ═════════════════════════════════════════════════════════════
# Eyelid TBIL Ternary (optimized models)
# ═════════════════════════════════════════════════════════════
print('\n' + '='*60)
print('  Eyelid TBIL Ternary (Optimized)')
print('='*60)
clean_mdf = pd.read_csv(os.path.join(BASE, 'data', 'clean_dataset_manifest.csv'))
tmap = {'mild': 0, 'moderate': 1, 'severe': 2}
j_df = clean_mdf[clean_mdf['category'].isin(['mild','moderate','severe']) & (clean_mdf['n_eyelid']>0)]

tbil_val = []
for _, r in j_df.iterrows():
    imgs = json.loads(r['eyelid_images'])
    tbil_val.append({'eyelids': imgs, 'faces': [], 'label': tmap[r['category']], 'pid': r['patient_id']})

random.seed(SEED)
by_l = {}
for p in tbil_val: by_l.setdefault(p['label'], []).append(p)
va_t = []
for l, g in by_l.items():
    random.shuffle(g); va_t.extend(g[:max(1, int(len(g)*0.2))])

print(f'  Val: {len(va_t)} {Counter(p["label"] for p in va_t)}')

EYE_TBIL = [
    ('TBIL-Eyelid-ConvNeXt', 'convnext_tiny', 'eyelid_opt_convnext_tiny.pt'),
    ('TBIL-Eyelid-ViT', 'vit_tiny_patch16_224', 'eyelid_opt_vit_tiny_patch16_224.pt'),
    ('TBIL-Eyelid-Swin', 'swin_tiny_patch4_window7_224', 'eyelid_opt_swin_tiny_patch4_window7_224.pt'),
]
tbil_preds, tbil_true = {}, None
for name, bb, ckpt in EYE_TBIL:
    r = get_preds_and_ci(name, bb, ckpt, 3, va_t, 'eyelids')
    if r:
        res, probs, true = r; res['task'] = 'TBIL-Eyelid'; res['modality'] = 'Eyelid'
        all_results.append(res); tbil_preds[name] = probs; tbil_true = true

if len(tbil_preds) >= 2 and tbil_true is not None:
    ens = np.mean(list(tbil_preds.values()), axis=0)
    bs = bootstrap_metrics(tbil_true, ens, 3)
    pa = roc_auc_score(tbil_true, ens, multi_class='ovr', labels=[0,1,2])
    all_results.append({'name': 'TBIL-Eyelid-Ensemble', 'task': 'TBIL-Eyelid', 'modality': 'Eyelid',
                         'nc': 3, 'n_val': len(tbil_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(tbil_true, ens.argmax(1)),
                         'point_f1': f1_score(tbil_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    TBIL-Eyelid-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

# ═════════════════════════════════════════════════════════════
# Jaundice Type (Hepato vs Chol)
# ═════════════════════════════════════════════════════════════
print('\n' + '='*60)
print('  Jaundice Type (Hepato vs Chol)')
print('='*60)
EXCEL = os.path.join(BASE, '213008980_按序号_肝炎患者补充问卷-医护_313_313.xlsx')
df_ex = pd.read_excel(EXCEL)
jt_col = [c for c in df_ex.columns if '黄疸类型' in str(c)][0]
df_ex['hid_str'] = df_ex['2、患者住院号'].astype(str)

def get_jt(pid):
    nums = re.findall(r'\d+', pid.replace('.zip', ''))
    hid = nums[-1] if nums else None
    if not hid: return None
    hn = hid.lstrip('0')
    rows = df_ex[df_ex['hid_str'].apply(lambda s: s.lstrip('0') if s else s) == hn]
    if len(rows) == 0:
        rows = df_ex[df_ex['hid_str'].apply(lambda s: s.lstrip('0') if s else s).str.endswith(hn[-6:])]
    return float(rows.iloc[0][jt_col]) if len(rows) > 0 else None

type_pats = {}
DATA = os.path.join(BASE, 'data', 'sam_processed_v3')
for cat in ['mild','moderate','severe']:
    cd = os.path.join(DATA, cat)
    for pid in os.listdir(cd):
        p = os.path.join(cd, pid)
        if not os.path.isdir(p): continue
        faces = sorted([os.path.join(p,f) for f in os.listdir(p) if '_face' in f])
        if not faces: continue
        jt = get_jt(pid)
        if jt in (1, 2): label = 0
        elif jt == 3: label = 1
        else: continue
        type_pats[pid] = {'faces': faces, 'label': label}

by_g = {}
for pid, p in type_pats.items(): by_g.setdefault(p['label'], []).append(pid)
va_type_pids = set()
random.seed(SEED)
for g, pids in by_g.items():
    random.shuffle(pids); va_type_pids.update(pids[:max(1, int(len(pids)*0.2))])
va_type = [{'faces': type_pats[pid]['faces'], 'label': type_pats[pid]['label'], 'pid': pid}
            for pid in va_type_pids]
print(f'  Val: {len(va_type)} {Counter(p["label"] for p in va_type)}')

TYPE_MODELS = [
    ('Type-EfficientNet', 'efficientnet_b0', 'type_class_efficientnet_b0.pt'),
    ('Type-Swin', 'swin_tiny_patch4_window7_224', 'type_class_swin_tiny_patch4_window7_224.pt'),
    ('Type-ViT', 'vit_tiny_patch16_224', 'type_class_vit_tiny_patch16_224.pt'),
    ('Type-ConvNeXt', 'convnext_tiny', 'type_class_convnext_tiny.pt'),
]
type_preds, type_true = {}, None
for name, bb, ckpt in TYPE_MODELS:
    r = get_preds_and_ci(name, bb, ckpt, 2, va_type, 'faces')
    if r:
        res, probs, true = r; res['task'] = 'Jaundice-Type'; res['modality'] = 'Face'
        all_results.append(res); type_preds[name] = probs; type_true = true

if len(type_preds) >= 2 and type_true is not None:
    ens = np.mean(list(type_preds.values()), axis=0)
    bs = bootstrap_metrics(type_true, ens, 2)
    pa = roc_auc_score(type_true, ens[:, 1])
    all_results.append({'name': 'Type-Ensemble', 'task': 'Jaundice-Type', 'modality': 'Face',
                         'nc': 2, 'n_val': len(type_true), 'point_auc': pa,
                         'auc_lo': bs['auc_lo'], 'auc_hi': bs['auc_hi'],
                         'point_acc': accuracy_score(type_true, ens.argmax(1)),
                         'point_f1': f1_score(type_true, ens.argmax(1), average='macro', zero_division=0)})
    print(f'    Type-Ensemble: AUC={pa:.4f} [{bs["auc_lo"]:.4f}-{bs["auc_hi"]:.4f}]')

# ═════════════════════════════════════════════════════════════
# SUMMARY + SAVE
# ═════════════════════════════════════════════════════════════
print('\n' + '='*80)
print(f'  BOOTSTRAP 95% CI SUMMARY ({N_BOOT} iterations, patient-level)')
print('='*80)
print(f'\n{"Task":<16} {"Model":<32} {"AUC [95% CI]":>26} {"N":>4}')
print('-' * 82)
for r in sorted(all_results, key=lambda x: (x.get('task',''), -x.get('point_auc',0))):
    print(f'{r.get("task",""):<16} {r["name"]:<32} {r["point_auc"]:.3f} [{r["auc_lo"]:.3f}-{r["auc_hi"]:.3f}]   {r.get("n_val",0):>4}')

pd.DataFrame(all_results).to_csv(os.path.join(TBL, 'bootstrap_ci_all_models.csv'), index=False, encoding='utf-8-sig')
with open(os.path.join(RES, 'bootstrap_ci_all_models.json'), 'w') as f:
    json.dump(all_results, f, indent=2, default=str)

# DOCX
doc = Document()
doc.add_heading('All Models with Bootstrap 95% Confidence Intervals', level=2)
doc.add_paragraph(f'Bootstrap: {N_BOOT} iterations, patient-level resampling. CI = [2.5th, 97.5th] percentile.')
for task in sorted(set(r.get('task','') for r in all_results)):
    sub = sorted([r for r in all_results if r.get('task') == task], key=lambda x: -x.get('point_auc',0))
    doc.add_heading(task, level=3)
    t = doc.add_table(rows=1, cols=5)
    t.style = 'Light Grid Accent 1'
    for i, c in enumerate(['Model', 'AUC', '95% CI Low', '95% CI High', 'N']):
        t.rows[0].cells[i].text = c
        for p in t.rows[0].cells[i].paragraphs:
            for run in p.runs: run.bold = True; run.font.size = Pt(8)
    for r in sub:
        cells = t.add_row().cells
        cells[0].text = r['name']
        cells[1].text = f'{r["point_auc"]:.4f}'
        cells[2].text = f'{r["auc_lo"]:.4f}'
        cells[3].text = f'{r["auc_hi"]:.4f}'
        cells[4].text = str(r.get('n_val', 0))
        for c in cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size = Pt(8)
doc.save(os.path.join(TBL, 'Table_bootstrap_ci_all_models.docx'))

# Forest plot
print('\nGenerating forest plot...')
tasks_list = sorted(set(r.get('task','') for r in all_results))
n_tasks = len(tasks_list)
fig, axes = plt.subplots(1, n_tasks, figsize=(4.5*n_tasks, max(6, len(all_results)*0.4/n_tasks)), sharey=False)
if n_tasks == 1: axes = [axes]
colors = {'DBIL': '#E76F51', 'IBIL': '#2A9D8F', 'TBIL-Eyelid': '#457B9D', 'Jaundice-Type': '#6A4C93'}
for ax, task in zip(axes, tasks_list):
    sub = sorted([r for r in all_results if r.get('task') == task], key=lambda x: x.get('point_auc',0))
    for i, r in enumerate(sub):
        ax.errorbar(r['point_auc'], i,
                     xerr=[[r['point_auc']-r['auc_lo']], [r['auc_hi']-r['point_auc']]],
                     fmt='o', color=colors.get(task, 'gray'), capsize=3, markersize=5, lw=1.2)
    ax.set_yticks(range(len(sub)))
    ax.set_yticklabels([r['name'].replace(f'{task}-','').replace('Eyelid','Eye') for r in sub], fontsize=7)
    ax.set_xlabel('AUC-ROC [95% CI]', fontsize=9)
    ax.set_title(task, fontsize=11, fontweight='bold')
    ax.axvline(0.5, color='gray', ls='--', lw=0.5, alpha=0.5)
    ax.set_xlim([0.3, 1.0]); ax.grid(axis='x', alpha=0.2)
plt.suptitle(f'Bootstrap 95% CI ({N_BOOT} iterations, patient-level)', fontsize=12, fontweight='bold', y=1.01)
plt.tight_layout()
fp = os.path.join(FIG, 'Fig_bootstrap_forest.png')
fig.savefig(fp); fig.savefig(fp.replace('.png','.svg'))
plt.close(fig)

print(f'\n  CSV:  {TBL}/bootstrap_ci_all_models.csv')
print(f'  DOCX: {TBL}/Table_bootstrap_ci_all_models.docx')
print(f'  Figure: {fp}')
