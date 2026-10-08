# -*- coding: utf-8 -*-
"""
Rebuild supplementary docx: insert ALL 13 figures at correct positions
by matching section headings, not just caption text.
"""
import os, re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r'D:\research\人脸识别营养\传染科'
SRC_MD = os.path.join(BASE, 'results', 'face_v4_manuscript_delivery', 'manuscript', 'supplementary_figures', 'Supplementary_Materials_v3.md')
FIG_DIR = os.path.join(BASE, 'results', 'face_v4_manuscript_delivery', 'manuscript', 'supplementary_figures')
OUT_DOCX = os.path.join(BASE, 'results', 'BilinGuard_Lancet_v2_20260725', '04_Supplementary', 'BilinGuard_Supplementary_Lancet_expanded.docx')

text = open(SRC_MD, encoding='utf-8').read()

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)

# Define: (section_heading_keyword, figure_filename, caption_text, width)
FIGURE_INSERTIONS = [
    ('3. Video Preprocessing', 'SuppFigS7_preprocessing_pipeline.png',
     'Supplementary Figure S7. Video preprocessing pipeline (v4: YuNet + SAM). 12-frame sampling, '
     'YuNet detection, SAM masking, skin-fraction gate, landmark alignment, CLAHE, 224x224 output.', 6.0),
    ('3.6 Grey-World', 'SuppFigS9_domain_adaptation_workflow.png',
     'Supplementary Figure S9. External cohort evaluation workflow. v4-lite preprocessing was selected '
     'after comparative evaluation on the same cohort.', 6.0),
    ('4. Layer 1', 'SuppFigS8_dataflow.png',
     'Supplementary Figure S8. Study design and data flow architecture (5 modules: cohort, '
     'preprocessing, evaluation, tasks, CDSS).', 6.5),
    ('Unified 105-cell', 'SuppFig1_forest_all_models.png',
     'Supplementary Figure S1. Forest plot of all model AUCs with 95% bootstrap CIs '
     '(canonical-v2, 105 evaluation cells).', 5.5),
    ('Key Performance Summary', 'SuppFig2_summary_3layer.png',
     'Supplementary Figure S2. Task-level performance summary (Swin-Tiny or ensemble, '
     'internal validation + external cohort).', 6.0),
    ('Dual-view late fusion', 'SuppFig5_binary_models.png',
     'Supplementary Figure S5. Binary screening model comparison (4 backbones + ensemble, '
     'ROC and precision-recall curves).', 6.0),
    ('5.5 Grad-CAM', 'SuppFig4_confusion_matrices.png',
     'Supplementary Figure S4. Confusion matrices for the best face model per task '
     '(canonical-v2, patient-level).', 6.0),
    ('6. Layer 3', 'SuppFigS10_cdss_decision_tree.png',
     'Supplementary Figure S10. CDSS decision flow (7 sequential steps; safety gate '
     'overrides all downstream logic).', 5.5),
    ('7. External', 'SuppFig3_domain_adaptation.png',
     'Supplementary Figure S3. Internal vs external probability distributions and calibration '
     '(Swin-Tiny). External slope = 3.10, Brier = 0.185.', 6.0),
    ('8. Training', 'SuppFigS12_training_procedure.png',
     'Supplementary Figure S12. Training and evaluation procedure (canonical-v2 protocol: '
     'patient-level split, augmentation, Focal Loss, EMA, bootstrap CI).', 6.0),
    ('9.6 Shared-Split', 'SuppFig6_grading_models.png',
     'Supplementary Figure S6. Ternary grading model comparison (face and eyelid ROC curves).', 6.0),
    ('10. Reader Study', 'SuppFigS11_reader_study_protocol.png',
     'Supplementary Figure S11. Reader study protocol. Eight human readers, blinded 4-class '
     'review of 100 cases.', 6.0),
    ('11. Desktop', 'SuppFigS13_desktop_architecture.png',
     'Supplementary Figure S13. Desktop application architecture (PyQt5 v4.0: UI, inference '
     'engine, CDSS module).', 6.0),
]

def add_figure(doc, fname, caption, width):
    img_path = os.path.join(FIG_DIR, fname)
    if not os.path.exists(img_path):
        print(f'  WARNING: {fname} not found')
        return False
    # Caption
    cap_p = doc.add_paragraph()
    cap_p.paragraph_format.space_before = Pt(8)
    cap_p.paragraph_format.space_after = Pt(4)
    run = cap_p.add_run(caption)
    run.italic = True
    run.font.size = Pt(8)
    # Image
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(img_path, width=Inches(width))
    # Spacer
    doc.add_paragraph('')
    return True

# Parse markdown and build document
lines = text.split('\n')
fig_count = 0
pending_figs = list(FIGURE_INSERTIONS)

for line in lines:
    s = line.rstrip()
    
    if not s.strip():
        doc.add_paragraph('')
        continue
    
    # Check if this line matches a figure insertion point
    s_lower = s.lower()
    matched = None
    for i, (heading_kw, fname, caption, width) in enumerate(pending_figs):
        if heading_kw.lower() in s_lower:
            matched = i
            break
    
    # Add the text
    if s.startswith('# '):
        doc.add_heading(s[2:].strip(), level=1)
    elif s.startswith('## '):
        doc.add_heading(s[3:].strip(), level=2)
    elif s.startswith('### '):
        doc.add_heading(s[4:].strip(), level=3)
    elif s.startswith('---'):
        doc.add_paragraph('-' * 40)
    elif s.startswith('> '):
        p = doc.add_paragraph()
        run = p.add_run(s[2:].strip()); run.italic = True
    else:
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*[^*]+\*\*)', s)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2]); run.bold = True
            else:
                p.add_run(part)
    
    # Insert figure after this paragraph if matched
    if matched is not None:
        heading_kw, fname, caption, width = pending_figs.pop(matched)
        if add_figure(doc, fname, caption, width):
            fig_count += 1
            print(f'  Inserted: {fname} (after "{heading_kw}")')

# Insert any remaining figures at the end
for heading_kw, fname, caption, width in pending_figs:
    if add_figure(doc, fname, caption, width):
        fig_count += 1
        print(f'  Inserted (end): {fname}')

doc.save(OUT_DOCX)

# Verify
img_total = 0
for p in doc.paragraphs:
    blips = p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    img_total += len(blips)
fsize = os.path.getsize(OUT_DOCX) / 1024 / 1024

print(f'\nDocument rebuilt: {fig_count} unique figures, {img_total} images total')
print(f'File size: {fsize:.1f} MB')
print(f'Paragraphs: {len(doc.paragraphs)}')
