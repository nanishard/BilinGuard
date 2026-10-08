# -*- coding: utf-8 -*-
"""Insert SuppFig1 and SuppFig2 with new captions."""
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = r'D:\research\人脸识别营养\传染科\results\BilinGuard_Lancet_v2_20260725\04_Supplementary\BilinGuard_Supplementary_Lancet_expanded.docx'
FIG_DIR = r'D:\research\人脸识别营养\传染科\results\BilinGuard_Lancet_v2_20260725\04_Supplementary\figures'

doc = Document(DOCX)

# Find insertion points
# SuppFig1 (forest plot): insert after "105-cell evaluation matrix" or "Key Performance Summary"
# SuppFig2 (task summary): insert after SuppFig1

target_idx = -1
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip().lower()
    if 'key performance summary' in t or '105-cell' in t or 'canonical-v2' in t and 'summary' in t:
        target_idx = i
        break
if target_idx < 0:
    # Try finding the performance summary table area
    for i, p in enumerate(doc.paragraphs):
        if 'performance' in p.text.lower() and 'summary' in p.text.lower():
            target_idx = i
            break

if target_idx < 0:
    # Insert before "Supplementary Figures" section
    for i, p in enumerate(doc.paragraphs):
        if 'supplementary figures' in p.text.lower() or 'figure s7' in p.text.lower():
            target_idx = i - 1
            break

print(f'Target paragraph for SuppFig1/2: {target_idx}')
if target_idx >= 0:
    print(f'  Text: {doc.paragraphs[target_idx].text[:80]}')

if target_idx >= 0:
    p_target = doc.paragraphs[target_idx]

    # Add SuppFig1 caption + image
    cap1 = doc.add_paragraph()
    cap1.paragraph_format.space_before = Pt(12)
    run1 = cap1.add_run('Supplementary Figure S1. Forest plot of all model AUCs with 95% bootstrap CIs '
                         '(canonical-v2, 7 tasks x 3 input configurations x 5 model variants = 105 cells).')
    run1.italic = True
    run1.font.size = Pt(8)

    img1 = doc.add_paragraph()
    img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img1.add_run().add_picture(os.path.join(FIG_DIR, 'SuppFig1_forest_all_models.png'), width=Inches(6.0))

    # Add SuppFig2 caption + image
    cap2 = doc.add_paragraph()
    cap2.paragraph_format.space_before = Pt(12)
    run2 = cap2.add_run('Supplementary Figure S2. Task-level performance summary. Best AUC per task '
                         '(Swin-Tiny or four-backbone ensemble, internal validation). External cohort '
                         'evaluation (Swin, v4-lite preprocessing) shown in terracotta.')
    run2.italic = True
    run2.font.size = Pt(8)

    img2 = doc.add_paragraph()
    img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img2.add_run().add_picture(os.path.join(FIG_DIR, 'SuppFig2_summary_3layer.png'), width=Inches(6.0))

    # Move all 4 new paragraphs to after target
    p_element = p_target._element
    for new_p in [cap2._element, img2._element, cap1._element, img1._element]:
        p_element.addnext(new_p)

    doc.save(DOCX)
    print('SuppFig1 + SuppFig2 inserted with captions')
else:
    # Fallback: add at end of document
    doc.add_paragraph()
    cap1 = doc.add_paragraph()
    run1 = cap1.add_run('Supplementary Figure S1. Forest plot of all model AUCs (canonical-v2).')
    run1.italic = True; run1.font.size = Pt(8)
    img1 = doc.add_paragraph()
    img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img1.add_run().add_picture(os.path.join(FIG_DIR, 'SuppFig1_forest_all_models.png'), width=Inches(6.0))

    cap2 = doc.add_paragraph()
    run2 = cap2.add_run('Supplementary Figure S2. Task-level performance summary.')
    run2.italic = True; run2.font.size = Pt(8)
    img2 = doc.add_paragraph()
    img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img2.add_run().add_picture(os.path.join(FIG_DIR, 'SuppFig2_summary_3layer.png'), width=Inches(6.0))

    doc.save(DOCX)
    print('SuppFig1 + SuppFig2 appended at end')

# Count total images
from docx.opc.constants import RELATIONSHIP_TYPE as RT
img_count = sum(1 for rel in doc.part.rels.values() if 'image' in rel.reltype)
print(f'Total images in document: {img_count}')
