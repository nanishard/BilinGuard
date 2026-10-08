# -*- coding: utf-8 -*-
"""
Rebuild BilinGuard_Supplementary_Lancet_expanded.docx from scratch.
Read Supplementary_Materials_v3.md, create clean docx with each figure inserted exactly once.
"""
import os, re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r'D:\research\人脸识别营养\传染科'
SRC_MD = os.path.join(BASE, 'results', 'face_v4_manuscript_delivery', 'manuscript', 'supplementary_figures', 'Supplementary_Materials_v3.md')
FIG_DIR = os.path.join(BASE, 'results', 'face_v4_manuscript_delivery', 'manuscript', 'supplementary_figures')
OUT_DOCX = os.path.join(BASE, 'results', 'BilinGuard_Lancet_v2_20260725', '04_Supplementary', 'BilinGuard_Supplementary_Lancet_expanded.docx')

# Figure mapping: keyword in caption -> (filename, width_inches)
FIG_MAP = {
    'forest plot': ('SuppFig1_forest_all_models.png', 6.0),
    'task-level performance': ('SuppFig2_summary_3layer.png', 6.0),
    'probability distribution': ('SuppFig3_domain_adaptation.png', 6.0),
    'calibration': ('SuppFig3_domain_adaptation.png', 6.0),
    'confusion matri': ('SuppFig4_confusion_matrices.png', 6.0),
    'binary screening model comparison': ('SuppFig5_binary_models.png', 6.0),
    'grading model': ('SuppFig6_grading_models.png', 6.0),
    'preprocessing pipeline': ('SuppFigS7_preprocessing_pipeline.png', 6.0),
    'data flow': ('SuppFigS8_dataflow.png', 6.5),
    'dataflow': ('SuppFigS8_dataflow.png', 6.5),
    'closed-loop': ('SuppFigS8_dataflow.png', 6.5),
    'domain adaptation workflow': ('SuppFigS9_domain_adaptation_workflow.png', 6.5),
    'external-cohort preprocessing': ('SuppFigS9_domain_adaptation_workflow.png', 6.5),
    'decision flow': ('SuppFigS10_cdss_decision_tree.png', 6.0),
    'decision tree': ('SuppFigS10_cdss_decision_tree.png', 6.0),
    'reader study protocol': ('SuppFigS11_reader_study_protocol.png', 6.5),
    'training procedure': ('SuppFigS12_training_procedure.png', 6.5),
    'desktop': ('SuppFigS13_desktop_architecture.png', 6.5),
}

# Read markdown
text = open(SRC_MD, encoding='utf-8').read()

# Create fresh document
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)

lines = text.split('\n')
inserted_figs = set()
fig_count = 0

for line in lines:
    s = line.rstrip()
    
    if not s.strip():
        doc.add_paragraph('')
        continue
    
    # Headings
    if s.startswith('# '):
        doc.add_heading(s[2:].strip(), level=1)
    elif s.startswith('## '):
        doc.add_heading(s[3:].strip(), level=2)
    elif s.startswith('### '):
        doc.add_heading(s[4:].strip(), level=3)
    elif s.startswith('---'):
        doc.add_paragraph('-' * 40)
    elif s.startswith('> '):
        p = doc.add_paragraph()
        run = p.add_run(s[2:].strip())
        run.italic = True
    else:
        # Check if this line is a figure caption
        is_caption = ('supplementary figure' in s.lower() or 
                      'suppfig' in s.lower() or
                      'figure s' in s.lower())
        
        # Add the text paragraph
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*[^*]+\*\*)', s)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        
        # If it's a figure caption, try to insert the image after it
        if is_caption:
            s_lower = s.lower()
            matched_fig = None
            for keyword, (fname, width) in FIG_MAP.items():
                if keyword in s_lower and fname not in inserted_figs:
                    img_path = os.path.join(FIG_DIR, fname)
                    if os.path.exists(img_path):
                        matched_fig = (fname, width)
                        break
            
            if matched_fig:
                fname, width = matched_fig
                inserted_figs.add(fname)
                # Add image paragraph
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(os.path.join(FIG_DIR, fname), width=Inches(width))
                fig_count += 1
                print(f'  Inserted: {fname} (after caption at para)')

doc.save(OUT_DOCX)

# Verify
img_total = sum(1 for p in doc.paragraphs 
                for _ in p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))
fsize = os.path.getsize(OUT_DOCX) / 1024 / 1024

print(f'\nDocument rebuilt: {fig_count} figures inserted (no duplicates)')
print(f'File: {OUT_DOCX}')
print(f'Size: {fsize:.1f} MB')
