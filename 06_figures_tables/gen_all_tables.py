# -*- coding: utf-8 -*-
"""
Generate all manuscript tables (Table 1-5) + supplementary tables.

Outputs:
  results/tables_v3/Table1_baseline.csv
  results/tables_v3/Table2_layer1_recognition.csv
  results/tables_v3/Table3_layer2_scoring.csv
  results/tables_v3/Table4_reader_study.csv
  results/tables_v3/Table5_cdss_knowledge.csv
  results/tables_v3/SuppTable_S1_all_67_models.csv
  results/tables_v3/SuppTable_S2_64_guidelines.csv
  results/tables_v3/ALL_TABLES.docx  (consolidated Word document)
"""
import os, sys, json, re, pandas as pd, numpy as np
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r'D:\research\人脸识别营养\传染科'
RES = os.path.join(BASE, 'results')
TBL_V3 = os.path.join(RES, 'tables_v3')
sys.path.insert(0, os.path.join(BASE, 'deployment'))

bs = pd.read_csv(os.path.join(TBL_V3, 'bootstrap_full_metrics_v3.csv'))

def fmt(v, lo, hi):
    if pd.isna(v): return '-'
    if pd.isna(lo) or pd.isna(hi): return f'{v:.3f}'
    return f'{v:.3f} [{lo:.3f}-{hi:.3f}]'

# ═════════════════════════════════════════════════════════════
# Table 1: Baseline Characteristics
# ═════════════════════════════════════════════════════════════
print('[1] Table 1: Baseline Characteristics...')
# Use existing baseline data + corrected counts
try:
    df_base = pd.read_excel(os.path.join(BASE, 'baseline6.17.xlsx'))
    df_base.columns = ['name', 'hid', 'diagnosis', 'group', 'gender', 'age', 'height', 'weight',
                        'bmi', 'heart_rate', 'waist', 'hip'][:len(df_base.columns)]
    g1 = df_base[df_base['group'] == 1]  # non-jaundice
    g2 = df_base[df_base['group'] == 2]  # jaundice

    def mean_sd(series):
        s = pd.to_numeric(series, errors='coerce').dropna()
        return f'{s.mean():.2f} ± {s.std():.2f}' if len(s) > 0 else '-'

    def n_pct(series, val):
        n = (series == val).sum()
        return f'{n} ({n/len(series)*100:.1f})'

    table1 = pd.DataFrame({
        'Characteristic': [
            'N', 'Age (years), mean ± SD', 'Male sex, n (%)',
            'BMI (kg/m²), mean ± SD', 'Heart rate (bpm), mean ± SD',
            'Waist circumference (cm), mean ± SD',
            'Hip circumference (cm), mean ± SD',
        ],
        'Non-Jaundice Controls (n=625)': [
            '625', mean_sd(g1['age']), n_pct(g1['gender'], 1),
            mean_sd(g1['bmi']), mean_sd(g1['heart_rate']),
            mean_sd(g1['waist']), mean_sd(g1['hip']),
        ],
        'Jaundice Internal (n=313)': [
            '313', mean_sd(g2['age']), n_pct(g2['gender'], 1),
            mean_sd(g2['bmi']), mean_sd(g2['heart_rate']),
            mean_sd(g2['waist']), mean_sd(g2['hip']),
        ],
        'Jaundice External (n=58)': [
            '58', '59.54 ± 11.93', '31 (53.4)',
            '21.91 ± 2.75', '79.03 ± 18.91',
            '87.27 ± 7.20', '90.91 ± 13.68',
        ],
    })
    table1.to_csv(os.path.join(TBL_V3, 'Table1_baseline.csv'), index=False, encoding='utf-8-sig')
    print('  Saved Table1_baseline.csv')
except Exception as e:
    print(f'  ERROR: {e}')


# ═════════════════════════════════════════════════════════════
# Table 2: Layer 1 Recognition Performance
# ═════════════════════════════════════════════════════════════
print('[2] Table 2: Layer 1 Recognition Performance...')
l1_tasks = ['Binary-Screening', 'Binary-External-Pure', 'Binary-External',
             'TBIL-Eyelid', 'Face-Ternary', 'DBIL', 'IBIL', 'Jaundice-Type']
l1_df = bs[bs['task'].isin(l1_tasks)].copy()
# Sort by task then AUC descending
l1_df['_task_order'] = l1_df['task'].map({t: i for i, t in enumerate(l1_tasks)})
l1_df = l1_df.sort_values(['_task_order', 'auc'], ascending=[True, False])

rows = []
for _, r in l1_df.iterrows():
    rows.append({
        'Task': r['task'], 'Model': r['name'], 'N': int(r['n_val']),
        'F1 [95% CI]': fmt(r.get('f1', 0), r.get('f1_lo', 0), r.get('f1_hi', 0)),
        'Sensitivity [95% CI]': fmt(r.get('sens', 0), r.get('sens_lo', 0), r.get('sens_hi', 0)),
        'Specificity [95% CI]': fmt(r.get('spec', 0), r.get('spec_lo', 0), r.get('spec_hi', 0)),
        'Accuracy [95% CI]': fmt(r.get('acc', 0), r.get('acc_lo', 0), r.get('acc_hi', 0)),
        'AUC-ROC [95% CI]': fmt(r['auc'], r['auc_lo'], r['auc_hi']),
        'AP [95% CI]': fmt(r.get('ap', 0), r.get('ap_lo', 0), r.get('ap_hi', 0)),
    })
table2 = pd.DataFrame(rows)
table2.to_csv(os.path.join(TBL_V3, 'Table2_layer1_recognition.csv'), index=False, encoding='utf-8-sig')
print(f'  Saved Table2 ({len(table2)} rows)')


# ═════════════════════════════════════════════════════════════
# Table 3: Layer 2 Prognostic Scoring
# ═════════════════════════════════════════════════════════════
print('[3] Table 3: Layer 2 Prognostic Scoring...')
l2_tasks = ['Child-Pugh', 'MELD-Risk']
l2_df = bs[bs['task'].isin(l2_tasks)].sort_values(['task', 'auc'], ascending=[True, False])
rows = []
for _, r in l2_df.iterrows():
    rows.append({
        'Task': r['task'].replace('Child-Pugh', 'Child-Pugh (A/B/C)').replace('MELD-Risk', 'MELD Risk (Low/Med/High)'),
        'Model': r['name'], 'N': int(r['n_val']),
        'F1 [95% CI]': fmt(r.get('f1', 0), r.get('f1_lo', 0), r.get('f1_hi', 0)),
        'Sensitivity [95% CI]': fmt(r.get('sens', 0), r.get('sens_lo', 0), r.get('sens_hi', 0)),
        'Specificity [95% CI]': fmt(r.get('spec', 0), r.get('spec_lo', 0), r.get('spec_hi', 0)),
        'Accuracy [95% CI]': fmt(r.get('acc', 0), r.get('acc_lo', 0), r.get('acc_hi', 0)),
        'AUC-ROC [95% CI]': fmt(r['auc'], r['auc_lo'], r['auc_hi']),
        'AP [95% CI]': fmt(r.get('ap', 0), r.get('ap_lo', 0), r.get('ap_hi', 0)),
    })
table3 = pd.DataFrame(rows)
table3.to_csv(os.path.join(TBL_V3, 'Table3_layer2_scoring.csv'), index=False, encoding='utf-8-sig')
print(f'  Saved Table3 ({len(table3)} rows)')


# ═════════════════════════════════════════════════════════════
# Table 4: Reader Study
# ═════════════════════════════════════════════════════════════
print('[4] Table 4: Reader Study...')
try:
    reader = json.load(open(os.path.join(RES, 'reader_study_results.json'), encoding='utf-8'))
    doc_df = pd.read_csv(os.path.join(RES, 'tables', 'doctor_individual_metrics.csv'))

    # Per-rater binary
    rows = []
    for _, r in doc_df[doc_df['task'] == 'Binary'].iterrows():
        rows.append({
            'Rater': r['doctor'], 'Task': 'Binary Screening',
            'AUC': f'{r["auc"]:.3f}', 'F1': f'{r["f1"]:.3f}',
            'Sensitivity': f'{r["sens"]:.3f}', 'Specificity': f'{r["spec"]:.3f}',
            'Accuracy': f'{r["acc"]:.3f}', 'N': int(r['n']),
        })
    # BilinGuard binary
    rows.append({'Rater': 'BilinGuard', 'Task': 'Binary Screening',
                  'AUC': '0.993', 'F1': '0.933', 'Sensitivity': '0.976',
                  'Specificity': '0.913', 'Accuracy': '0.936', 'N': 110})
    # Per-rater ternary
    for _, r in doc_df[doc_df['task'] == 'Ternary'].iterrows():
        rows.append({
            'Rater': r['doctor'], 'Task': 'Three-Class Grading',
            'AUC': f'{r["auc"]:.3f}', 'F1': f'{r["f1"]:.3f}',
            'Sensitivity': f'{r["sens"]:.3f}', 'Specificity': f'{r["spec"]:.3f}',
            'Accuracy': f'{r["acc"]:.3f}', 'N': int(r['n']),
        })
    # BilinGuard ternary
    rows.append({'Rater': 'BilinGuard', 'Task': 'Three-Class Grading',
                  'AUC': '0.798', 'F1': '0.550', 'Sensitivity': '0.565',
                  'Specificity': '0.803', 'Accuracy': '0.715', 'N': 110})

    table4 = pd.DataFrame(rows)
    table4.to_csv(os.path.join(TBL_V3, 'Table4_reader_study.csv'), index=False, encoding='utf-8-sig')

    # Add summary statistics
    summary_rows = [
        {'Metric': "Inter-rater Fleiss' kappa (binary)", 'Value': f'{reader["inter_rater"]["fleiss_kappa"]:.2f}'},
        {'Metric': "Model-rater Cohen's kappa (mean±SD)", 'Value': f'{reader["model_rater"]["cohen_kappa_mean"]:.2f}±{reader["model_rater"]["cohen_kappa_std"]:.2f}'},
        {'Metric': 'Wilcoxon p-value range (model vs raters)', 'Value': f'{min(w["p_value"] for w in reader["wilcoxon"].values()):.2e} - {max(w["p_value"] for w in reader["wilcoxon"].values()):.3f}'},
        {'Metric': 'Time efficiency (median, Kruskal-Wallis)', 'Value': f'H={reader["time_efficiency"]["kruskal_wallis"]["H"]:.1f}, p={reader["time_efficiency"]["kruskal_wallis"]["p"]:.4f}'},
    ]
    pd.DataFrame(summary_rows).to_csv(os.path.join(TBL_V3, 'Table4_reader_study_stats.csv'),
                                        index=False, encoding='utf-8-sig')
    print(f'  Saved Table4 ({len(table4)} rows + stats)')
except Exception as e:
    print(f'  ERROR: {e}')


# ═════════════════════════════════════════════════════════════
# Table 5: CDSS Knowledge Base Summary
# ═════════════════════════════════════════════════════════════
print('[5] Table 5: CDSS Knowledge Base Summary...')
try:
    from clinical_advisor import (GUIDELINES, DEPARTMENTS, DISEASES, JAUNDICE_TYPES,
                                    SEVERITY_GUIDANCE, CLINICAL_PATHWAYS)

    # Guidelines by category
    guide_cats = {
        'Chinese Domestic': 0, 'International Core': 0, 'MELD/Child-Pugh Evidence': 0,
        'Expanded Hepatology': 0, 'Obstetrics/Pregnancy': 0, 'Nephrology': 0,
        'Nutrition/Other': 0,
    }
    for gid, info in GUIDELINES.items():
        source = info[3] if len(info) > 3 else ''
        if 'CSH' in source or 'CSID' in source or 'CSS' in source or 'NHC' in source or 'CSGE' in source or 'CMA' in source or 'CLTR' in source or 'CSOB' in source:
            guide_cats['Chinese Domestic'] += 1
        elif 'EASL' in source or 'AASLD' in source or 'APASL' in source or 'ACG' in source or 'Tokyo' in source or "King's" in source:
            guide_cats['International Core'] += 1
        elif 'MELD' in gid or 'CHILD' in gid or 'ALBI' in gid or 'OKUDA' in gid:
            guide_cats['MELD/Child-Pugh Evidence'] += 1
        elif 'EASL' in gid or 'AASLD' in gid:
            guide_cats['Expanded Hepatology'] += 1
        elif 'RCOG' in gid or 'ACOG' in gid or 'ICP' in gid or 'PREECLAMPSIA' in gid or 'HDP' in gid:
            guide_cats['Obstetrics/Pregnancy'] += 1
        elif 'KDIGO' in gid or 'ICA' in gid or 'TMA' in gid:
            guide_cats['Nephrology'] += 1
        else:
            guide_cats['Nutrition/Other'] += 1

    table5 = pd.DataFrame([
        {'Component': 'Clinical Guidelines (total)', 'Count': len(GUIDELINES),
         'Details': '12 Chinese + 52 International'},
        {'Component': '  Chinese Domestic', 'Count': guide_cats['Chinese Domestic'],
         'Details': 'CSH, CSID, CSS, NHC, CSGE guidelines'},
        {'Component': '  International Core', 'Count': guide_cats['International Core'],
         'Details': 'EASL, AASLD, APASL, ACG, Tokyo, King\'s College'},
        {'Component': '  MELD/Child-Pugh Evidence', 'Count': guide_cats['MELD/Child-Pugh Evidence'],
         'Details': 'MELD, MELD-Na, MELD 3.0, ALBI, Okuda, Child-Pugh, Child-Turcotte'},
        {'Component': 'Specialties/Departments', 'Count': len(DEPARTMENTS),
         'Details': 'Including department x jaundice-type decision matrix (51 rules)'},
        {'Component': 'Disease Profiles', 'Count': len(DISEASES),
         'Details': 'ALF, ACLF, cholangitis, HCC, hemolysis, PSC, PBC, AIH, DILI, etc.'},
        {'Component': 'Jaundice Types', 'Count': len(JAUNDICE_TYPES),
         'Details': 'Hepatocellular, Cholestatic, Hemolytic'},
        {'Component': 'Severity Levels', 'Count': len(SEVERITY_GUIDANCE),
         'Details': 'Levels 0-3 (TBIL thresholds: 34/85/171 umol/L)'},
        {'Component': 'Red-flag Clinical Pathways', 'Count': len(CLINICAL_PATHWAYS),
         'Details': 'ALF, ACLF, acute cholangitis, obstruction, hemolytic crisis, DILI'},
        {'Component': 'Laboratory Parameters', 'Count': 9,
         'Details': 'TBIL, DBIL, IBIL, ALT, AST, ALP, GGT, INR, Albumin'},
        {'Component': 'advise() Output Sections (max)', 'Count': 8,
         'Details': 'Severity, Type, Aetiology, CP/MELD, Dept, Alert, Labs, Triage'},
    ])
    table5.to_csv(os.path.join(TBL_V3, 'Table5_cdss_knowledge.csv'), index=False, encoding='utf-8-sig')
    print(f'  Saved Table5 ({len(table5)} rows)')
except Exception as e:
    print(f'  ERROR: {e}')


# ═════════════════════════════════════════════════════════════
# Supplementary Table S1: All 67 Models
# ═════════════════════════════════════════════════════════════
print('[S1] Supplementary Table S1: All 67 Models...')
supp_s1 = bs[['name', 'task', 'modality', 'n_val',
              'auc', 'auc_lo', 'auc_hi', 'f1', 'f1_lo', 'f1_hi',
              'sens', 'sens_lo', 'sens_hi', 'spec', 'spec_lo', 'spec_hi',
              'acc', 'acc_lo', 'acc_hi', 'ap', 'ap_lo', 'ap_hi']].copy()
supp_s1 = supp_s1.sort_values(['task', 'auc'], ascending=[True, False])
supp_s1.to_csv(os.path.join(TBL_V3, 'SuppTable_S1_all_67_models.csv'), index=False, encoding='utf-8-sig')
print(f'  Saved SuppTable_S1 ({len(supp_s1)} rows)')


# ═════════════════════════════════════════════════════════════
# Supplementary Table S2: 64 Guidelines
# ═════════════════════════════════════════════════════════════
print('[S2] Supplementary Table S2: 64 Guidelines...')
try:
    guide_rows = []
    for gid, info in GUIDELINES.items():
        guide_rows.append({
            'GID': gid,
            'Short Title': info[0],
            'Full Citation (EN)': info[1][:120] + '...' if len(info[1]) > 120 else info[1],
            'Source': info[3] if len(info) > 3 else '',
            'Year': info[4] if len(info) > 4 else '',
        })
    supp_s2 = pd.DataFrame(guide_rows).sort_values('Year', ascending=False)
    supp_s2.to_csv(os.path.join(TBL_V3, 'SuppTable_S2_64_guidelines.csv'), index=False, encoding='utf-8-sig')
    print(f'  Saved SuppTable_S2 ({len(supp_s2)} rows)')
except Exception as e:
    print(f'  ERROR: {e}')


# ═════════════════════════════════════════════════════════════
# Supplementary Table S3: 17 Diseases
# ═════════════════════════════════════════════════════════════
print('[S3] Supplementary Table S3: 17 Diseases...')
try:
    disease_rows = []
    for did, info in DISEASES.items():
        disease_rows.append({
            'Disease ID': did,
            'Name (EN)': info.get('name_en', ''),
            'Name (CN)': info.get('name_cn', ''),
            'Jaundice Type': info.get('jaundice_type', ''),
            'Priority': info.get('priority', ''),
            'Default Department': info.get('department', ''),
            'Red-flag Pathway': info.get('pathway', 'None'),
        })
    supp_s3 = pd.DataFrame(disease_rows).sort_values('Priority', ascending=False)
    supp_s3.to_csv(os.path.join(TBL_V3, 'SuppTable_S3_17_diseases.csv'), index=False, encoding='utf-8-sig')
    print(f'  Saved SuppTable_S3 ({len(supp_s3)} rows)')
except Exception as e:
    print(f'  ERROR: {e}')


# ═════════════════════════════════════════════════════════════
# Consolidated DOCX
# ═════════════════════════════════════════════════════════════
print('\n[DOCX] Creating consolidated Word document...')
doc = Document()
doc.add_heading('BilinGuard — Complete Results Tables', level=1)
doc.add_paragraph('67 models across 10 tasks. Bootstrap 1000 iterations, patient-level. '
                   '3-layer closed-loop architecture.')

tables_to_add = [
    ('Table 1. Baseline Characteristics', 'Table1_baseline.csv'),
    ('Table 2. Layer 1 — Jaundice Recognition Performance', 'Table2_layer1_recognition.csv'),
    ('Table 3. Layer 2 — Prognostic Scoring (Non-Invasive)', 'Table3_layer2_scoring.csv'),
    ('Table 4. Reader Study — BilinGuard vs Eight Clinicians', 'Table4_reader_study.csv'),
    ('Table 5. CDSS Knowledge Base Summary', 'Table5_cdss_knowledge.csv'),
    ('Supplementary Table S1. All 67 Models — Complete Metrics', 'SuppTable_S1_all_67_models.csv'),
    ('Supplementary Table S2. 64 Clinical Guidelines', 'SuppTable_S2_64_guidelines.csv'),
    ('Supplementary Table S3. 17 Disease Profiles', 'SuppTable_S3_17_diseases.csv'),
]
for title, fname in tables_to_add:
    path = os.path.join(TBL_V3, fname)
    if not os.path.exists(path):
        continue
    df = pd.read_csv(path)
    doc.add_heading(title, level=2)
    # Create table
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = 'Light Grid Accent 1'
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = str(c)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True; r.font.size = Pt(7)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7)

doc.save(os.path.join(TBL_V3, 'ALL_TABLES.docx'))
print(f'  Saved ALL_TABLES.docx')

print('\n' + '='*70)
print('  All tables generated.')
print('='*70)
print(f'\n  Output directory: {TBL_V3}')
for f in sorted(os.listdir(TBL_V3)):
    if f.endswith(('.csv', '.docx')):
        size = os.path.getsize(os.path.join(TBL_V3, f)) / 1024
        print(f'    {f:45s} {size:6.1f} KB')
