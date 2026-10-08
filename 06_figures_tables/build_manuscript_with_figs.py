# -*- coding: utf-8 -*-
"""
Build a new manuscript docx with all main-text figures and tables embedded,
each with its caption (figure legend / table title). Base text + styles are
preserved from the polished manuscript; images and rendered tables are inserted
immediately after their caption paragraphs.

Figure 5 caption is updated to match the new 8-panel figure (CDSS banner +
technical validation + six reader-study panels, internal+external ROCs).
"""
import os, csv, copy
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r'D:\research\人脸识别营养\传染科\results\BilinGuard_Lancet_v2_20260725'
SRC = os.path.join(BASE, '01_Manuscript', 'BilinGuard_Lancet_final_deep_polished.docx')
FIG = os.path.join(BASE, '02_Figures')
TAB = os.path.join(BASE, '03_Tables')
OUT = os.path.join(BASE, '01_Manuscript', 'BilinGuard_Lancet_with_figures_tables.docx')

FIG_FILES = {
    1: 'Figure1_system_overview.png',
    2: 'Figure2_binary_screening_v2.png',
    3: 'Figure3_grading_type_v2.png',
    4: 'Figure4_prognostic_v2.png',
    5: 'Figure5.png',
    6: 'Figure6_model_matrix.png',
}
TAB_FILES = {
    1: 'Table1_baseline.csv',
    2: 'Table2_layer1_trimmed.csv',
    3: 'Table3_layer2_hepatic_reserve.csv',
    4: 'Table4_reader_study.csv',
    5: 'Table5_cdss_validation.csv',
}

FIG5_CAPTION = (
    'Figure 5: Clinical decision support and reader study \u2014 BilinGuard versus '
    'eight clinicians. (A) CDSS architecture and rule-execution pathway (inputs, '
    'safety layer, rule engine, outputs; 64 guidelines, 51 rules, 6 red-flag '
    'pathways). (B) Knowledge base and technical validation; 150 of 150 prespecified '
    'synthetic scenarios passed. (C) Binary screening ROC: BilinGuard internal '
    '(n=112; AUC 0\u00b7956) and exploratory external (n=117; AUC 0\u00b7876) with 95% '
    'CIs, eight reader operating points, and reader mean (AUC 0\u00b7914). (D) '
    'Three-class grading macro-ROC: BilinGuard internal (n=42; AUC 0\u00b7880) and '
    'external (n=57; AUC 0\u00b7640) with 95% CIs, reader points, and reader mean '
    '(AUC 0\u00b7656). (E) Per-rater binary performance (AUC, accuracy, F1). (F) Binary '
    'sensitivity versus specificity. (G) Time efficiency per case. (H) Wilcoxon '
    'signed-rank comparison of BilinGuard versus each reader. AUC=area under the '
    'receiver operating characteristic curve. CDSS=clinical decision-support system.'
)


def set_par_text(par, text):
    """Replace paragraph text, keeping the first run's formatting + paragraph style."""
    for r in par.runs[1:]:
        r._element.getparent().remove(r._element)
    if par.runs:
        par.runs[0].text = text
    else:
        par.add_run(text)


def insert_image_after(doc, cap_par, img_path, width_in=6.3):
    doc.add_picture(img_path, width=Inches(width_in))
    pic_par = doc.paragraphs[-1]
    pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_par._element.addnext(pic_par._element)
    return pic_par


def insert_table_after(doc, cap_par, csv_path, font_pt=8):
    with open(csv_path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    nr, nc = len(rows), max(len(r) for r in rows)
    tbl = doc.add_table(rows=nr, cols=nc)
    tbl.style = 'Table Grid'
    tbl.autofit = True
    for i, row in enumerate(rows):
        for j in range(nc):
            val = row[j] if j < len(row) else ''
            cell = tbl.cell(i, j)
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_pt)
                    if i == 0:
                        r.font.bold = True
    cap_par._element.addnext(tbl._element)
    return tbl


doc = Document(SRC)

# locate caption paragraphs by text prefix
fig_cap = {}   # number -> paragraph
tab_cap = {}
for par in doc.paragraphs:
    t = par.text.strip()
    for n in range(1, 7):
        if t.startswith('Figure %d:' % n):
            fig_cap[n] = par
    for n in range(1, 6):
        if t.startswith('Table %d:' % n):
            tab_cap[n] = par

print('found figure captions:', sorted(fig_cap.keys()))
print('found table captions:', sorted(tab_cap.keys()))

# update Figure 5 caption to match the new 8-panel figure
if 5 in fig_cap:
    set_par_text(fig_cap[5], FIG5_CAPTION)
    print('updated Figure 5 caption (8-panel).')

# insert figures AFTER their captions (process descending so addnext chaining stays clean)
for n in sorted(fig_cap.keys(), reverse=True):
    img = os.path.join(FIG, FIG_FILES[n])
    if os.path.exists(img):
        insert_image_after(doc, fig_cap[n], img)
        print('  inserted Figure %d image' % n)
    else:
        print('  MISSING Figure %d image: %s' % (n, img))

# insert tables AFTER their titles
for n in sorted(tab_cap.keys(), reverse=True):
    csvp = os.path.join(TAB, TAB_FILES[n])
    if os.path.exists(csvp):
        insert_table_after(doc, tab_cap[n], csvp)
        print('  inserted Table %d' % n)
    else:
        print('  MISSING Table %d csv: %s' % (n, csvp))

doc.save(OUT)
print('\nSaved:', OUT)
print('Size: %.1f KB' % (os.path.getsize(OUT) / 1024))
