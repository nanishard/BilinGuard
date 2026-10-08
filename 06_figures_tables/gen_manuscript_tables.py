# -*- coding: utf-8 -*-
"""
Regenerate manuscript table files from the definitive 105-cell inventory
(canonical-v2, 2026-07-20). Baseline rows (YOLO/XGBoost/RF/SVM from other
experiments) are preserved verbatim; BilinGuard DL rows use v2 values+CI.

Outputs (CSV + DOCX, old files backed up as .bak):
  Table2_binary_screening.csv/.docx      (+ BilinGuard v2 rows)
  Table3_ternary_grading.csv/.docx       (all v2 ternary rows)
  bilirubin_classification_summary.csv   (v2 DBIL/IBIL)
  eyelid_face_combined_results.csv       (v2 ternary eyelid vs face, per-class sens)
  performance_summary.csv                (BilinGuard row -> v2)
"""
import os, json, shutil, numpy as np, pandas as pd
from sklearn.metrics import recall_score

BASE = r'D:\research\人脸识别营养\传染科'
RES = os.path.join(BASE, 'results')
TBL = os.path.join(RES, 'face_v4_manuscript_delivery', 'tables')
INV = pd.read_csv(os.path.join(TBL, 'MODEL_INVENTORY_complete.csv'))
PROBS = json.load(open(os.path.join(RES, 'canonical_v2_probs.json'), encoding='utf-8'))

from docx import Document
from docx.shared import Pt


def cell(task, scope, arch):
    r = INV[(INV['Task'] == task) & (INV['Scope'] == scope) & (INV['Architecture'] == arch)]
    return r.iloc[0] if len(r) else None


def fmt(v, lo, hi):
    if pd.isna(v): return 'N/A'
    return f'{v:.3f} ({lo:.3f}-{hi:.3f})'


def row_metrics(task, scope, arch):
    r = cell(task, scope, arch)
    return {'Accuracy': fmt(r['acc'], r['acc_lo'], r['acc_hi']),
            'F1-score': fmt(r['f1'], r['f1_lo'], r['f1_hi']),
            'Sensitivity': fmt(r['sens'], r['sens_lo'], r['sens_hi']),
            'Specificity': fmt(r['spec'], r['spec_lo'], r['spec_hi']),
            'AUC-ROC': fmt(r['auc'], r['auc_lo'], r['auc_hi']),
            'Avg. Precision': fmt(r['ap'], r['ap_lo'], r['ap_hi'])}


def backup(name):
    p = os.path.join(TBL, name)
    if os.path.exists(p) and not os.path.exists(p + '.bak'):
        shutil.copy(p, p + '.bak')


def save_docx(path, title, df):
    doc = Document()
    doc.add_heading(title, level=2)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = 'Light Grid Accent 1'
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
        for p in t.rows[0].cells[i].paragraphs:
            for run in p.runs: run.bold = True; run.font.size = Pt(8)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for run in p.runs: run.font.size = Pt(8)
    doc.save(path)


# ═══ Table2: Binary Screening ═══
backup('Table2_binary_screening.csv'); backup('Table2_binary_screening.docx')
t2 = pd.read_csv(os.path.join(TBL, 'Table2_binary_screening.csv.bak'))
new_rows = []
for label, scope, arch in [
    ('BilinGuard Face-ViT (best single)', 'Face', 'ViT'),
    ('BilinGuard Face-Ensemble', 'Face', 'Ensemble'),
    ('BilinGuard Eyelid-Ensemble', 'Eyelid', 'Ensemble'),
    ('BilinGuard Fusion-ConvNeXt (best fusion)', 'Fusion', 'ConvNeXt'),
    ('BilinGuard Fusion-Ensemble', 'Fusion', 'Ensemble')]:
    m = row_metrics('Binary Screening', scope, arch)
    new_rows.append({'Model': label, **m})
t2 = pd.concat([t2, pd.DataFrame(new_rows)], ignore_index=True)
t2.to_csv(os.path.join(TBL, 'Table2_binary_screening.csv'), index=False)
save_docx(os.path.join(TBL, 'Table2_binary_screening.docx'),
          'Table 2. Binary jaundice screening performance (95% CI). Baselines from prior experiments; BilinGuard rows = canonical-v2 internal validation (2026-07-20).', t2)
print('Table2 done:', t2.shape)

# ═══ Table3: Ternary Grading ═══
backup('Table3_ternary_grading.csv'); backup('Table3_ternary_grading.docx')
rows = []
for scope in ['Face', 'Eyelid', 'Fusion']:
    for arch in ['ConvNeXt', 'ViT', 'Swin', 'EfficientNet', 'Ensemble']:
        r = cell('Ternary Grading', scope, arch)
        rows.append({'Model': f'{scope}-{arch}',
                     'Accuracy': fmt(r['acc'], r['acc_lo'], r['acc_hi']),
                     'F1 (macro)': fmt(r['f1'], r['f1_lo'], r['f1_hi']),
                     'Sensitivity': fmt(r['sens'], r['sens_lo'], r['sens_hi']),
                     'Specificity': fmt(r['spec'], r['spec_lo'], r['spec_hi']),
                     'AUC-ROC (OvR)': fmt(r['auc'], r['auc_lo'], r['auc_hi'])})
t3 = pd.DataFrame(rows)
t3.to_csv(os.path.join(TBL, 'Table3_ternary_grading.csv'), index=False)
save_docx(os.path.join(TBL, 'Table3_ternary_grading.docx'),
          'Table 3. Ternary jaundice grading (mild/moderate/severe), canonical-v2 internal validation with 95% CI (2026-07-20).', t3)
print('Table3 done:', t3.shape)

# ═══ Bilirubin classification summary (DBIL/IBIL) ═══
backup('bilirubin_classification_summary.csv')
rows = []
for task, full, std in [('DBIL', 'Direct Bilirubin', '<=10 / 10-68 / >68 umol/L'),
                        ('IBIL', 'Indirect Bilirubin', '<=20 / 20-50 / >50 umol/L')]:
    for scope in ['Face', 'Eyelid', 'Fusion']:
        for arch in ['ConvNeXt', 'ViT', 'Swin', 'EfficientNet', 'Ensemble']:
            r = cell(task, scope, arch)
            rows.append({'Bilirubin': task, 'Grading Standard': f'{full} ({std})',
                         'Modality': scope, 'Model': arch,
                         'AUC': fmt(r['auc'], r['auc_lo'], r['auc_hi']),
                         'Accuracy': fmt(r['acc'], r['acc_lo'], r['acc_hi']),
                         'F1': fmt(r['f1'], r['f1_lo'], r['f1_hi'])})
bc = pd.DataFrame(rows)
bc.to_csv(os.path.join(TBL, 'bilirubin_classification_summary.csv'), index=False)
print('bilirubin_classification done:', bc.shape)

# ═══ eyelid vs face combined (ternary, per-class sens) ═══
backup('eyelid_face_combined_results.csv')
def per_class_sens(task, scope, arch_key, nc=3):
    d = PROBS.get(task, {}).get(scope, {}).get(arch_key)
    if not d: return (None,) * nc
    t = np.array([v[1] for v in d.values()])
    p = np.array([np.argmax(v[0]) for v in d.values()])
    return tuple(round(x, 3) for x in recall_score(t, p, labels=list(range(nc)), average=None, zero_division=0))

rows = []
for scope, label in [('eyelid', 'Eyelid'), ('face', 'Face')]:
    for arch_key, arch in [('convnext', 'ConvNeXt'), ('vit', 'ViT'), ('swin', 'Swin'),
                           ('efficientnet', 'EfficientNet'), ('__ensemble__', 'Ensemble')]:
        r = cell('Ternary Grading', label, arch)
        if r is None or pd.isna(r['auc']): continue
        s0, s1, s2 = per_class_sens('Ternary Grading', scope, arch_key)
        rows.append({'Model': f'{label}: {arch}',
                     'AUC-ROC (OvR)': fmt(r['auc'], r['auc_lo'], r['auc_hi']),
                     'Accuracy': fmt(r['acc'], r['acc_lo'], r['acc_hi']),
                     'F1 (macro)': fmt(r['f1'], r['f1_lo'], r['f1_hi']),
                     'Sens (Mild)': s0, 'Sens (Mod)': s1, 'Sens (Sev)': s2})
efc = pd.DataFrame(rows)
efc.to_csv(os.path.join(TBL, 'eyelid_face_combined_results.csv'), index=False)
print('eyelid_face_combined done:', efc.shape)

# ═══ performance_summary.csv (BilinGuard row -> v2) ═══
backup('performance_summary.csv')
ps = pd.read_csv(os.path.join(TBL, 'performance_summary.csv.bak'))
fe = cell('Binary Screening', 'Face', 'Ensemble')
fc = cell('Binary Screening', 'Fusion', 'ConvNeXt')
for i, r in ps.iterrows():
    if str(r['model']).strip() == 'BilinGuard':
        ps.at[i, 'auc_roc'] = round(fe['auc'], 3)
        ps.at[i, 'sensitivity'] = round(fe['sens'], 3)
        ps.at[i, 'specificity'] = round(fe['spec'], 3)
ps.loc[len(ps)] = {'model': 'BilinGuard-Fusion', 'auc_roc': round(fc['auc'], 3),
                   'sensitivity': round(fc['sens'], 3), 'specificity': round(fc['spec'], 3)}
ps.to_csv(os.path.join(TBL, 'performance_summary.csv'), index=False)
print('performance_summary done:', ps.shape)
print('\nAll tables regenerated.')
